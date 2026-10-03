#!/usr/bin/env python3
"""
notifybot_builder.py - Join step for the Parallel Orchestrator

Builds the flat file notifybot.csv by enriching each notification.csv row with:
  * application details from Managed_Segments (.xlsx or .csv), matched on
    appid (the file's "Application Id" column)
  * cluster details from inventory.csv, matched on the cluster name

This script does NOT collect anything from OpenShift - it only joins three
files that already exist. notification_inventory.py remains unchanged and is
the producer of notification.csv.

Base rows:      notification.csv        (one row per namespace)
Lookup by appid: Managed_Segments.xlsx or Managed_Segments.csv
                                        (key column: "Application Id")

Managed_Segments format is chosen by file extension:
  .xlsx / .xlsm  first worksheet, first row = header (needs openpyxl)
  .csv           first row = header; UTF-8 (BOM ok) or Windows-1252;
                 delimiter , ; tab or | detected automatically
Lookup by cluster: inventory.csv        (key column: "cluster")

Column handling
---------------
* Output headers are read from the files themselves (not hard-coded), so extra
  columns flow through automatically. Only the key column NAMES are assumed:
  "Application Id" in the spreadsheet and "cluster" in inventory.csv.
* The key columns are dropped from the appended sets (appid already carries the
  Application Id, cluster is already present from notification.csv).
* If an incoming column name collides with one already in the row, it is
  prefixed - "seg_" for spreadsheet columns, "inv_" for inventory columns -
  so no data is lost and no header is duplicated. (e.g. inventory's "env"
  becomes "inv_env", since notification.csv already has "env".)
* A missing lookup (appid is "NA"/"N/A"/blank, appid not found, or cluster not
  found) leaves those cells blank and keeps the row.
"""

import argparse
import csv
import logging
import re
import sys
import time
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Sequence, Tuple

# openpyxl is imported only when an .xlsx/.xlsm segments file is given,
# so CSV-only runs do not need it installed.

# Encodings tried, in order, for a Managed_Segments.csv file.
CSV_ENCODINGS = ("utf-8-sig", "cp1252")


OUTPUT_FILENAME = "notifybot.csv"

# Spreadsheet columns from which to pull an SOEID (2 letters + 5 digits) and
# turn it into an @citi.com id that is merged into the `ids` list. The owner
# and support-manager cells look like "Black, Tamara (TL68763)"; the SOEID
# cell is already just the token (e.g. "MP06424").
SOEID_SOURCE_COLUMNS = [
    "Primary Business/Information Owner",
    "Support Manager",
    "Application Manager SOEID",
]
# Exactly two letters followed by five digits, bounded so a 6+-digit run
# (e.g. AB123456) does not yield a false 5-digit match.
SOEID_RE = re.compile(r"\b[A-Za-z]{2}\d{5}\b")
ID_DOMAIN = "@citi.com"

# The merged-id column: read from notification.csv's "ids", emitted as "email".
IDS_INPUT_COL  = "ids"
IDS_OUTPUT_COL = "email"
IDS_SEPARATOR  = ";"    # separate each id with ';'
IDS_TRAILING   = False  # no trailing ';' after the last id (set True to add one)

# Key column names in the two lookup files.
SEGMENTS_KEY = "Application Id"
INVENTORY_KEY = "cluster"

# Prefixes applied to appended columns whose name collides with an existing one.
SEG_PREFIX = "seg_"
INV_PREFIX = "inv_"

# appid values that mean "no application" and should not be looked up.
EMPTY_APPIDS = {"", "na", "n/a", "none", "null"}

DEFAULT_OUTPUT_DIR = Path("/data/OSE/output")


def setup_logging(log_level: str = "INFO") -> logging.Logger:
    logging.basicConfig(
        level=getattr(logging, log_level),
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )
    return logging.getLogger("notifybot")


def squash_ws(value) -> str:
    """Trim the ends and collapse internal whitespace runs to one space.

    Used for JOIN KEYS and HEADER NAMES so that inconsistent spacing
    ("Application  Id", "gcbsg01p ", a tab vs a space) still matches. It is
    NOT applied to display values, so names like "Payments Hub" keep their
    real spacing.
    """
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_appid(value) -> str:
    """Normalize an appid / Application Id to a comparable string.

    Spreadsheet cells often come back as numbers (e.g. 123456 or 123456.0),
    while the appid from a namespace is a digit string. Collapse both to a
    plain integer string when the value is integer-valued, else str().strip().
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = squash_ws(value)
    # CSV cells are always text: "123456.0" (a number exported with decimals)
    # is the same Application Id as 123456.
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def soeid_emails_from(cell: str) -> List[str]:
    """Return @citi.com id(s) for every SOEID (2 letters + 5 digits) in `cell`.

    Works for both the "Name (XX#####)" owner/manager cells and the bare-token
    "Application Manager SOEID" cell. SOEIDs are lower-cased before the domain
    is appended (matching the manager-id convention in the original script).
    """
    if not cell:
        return []
    return [m.lower() + ID_DOMAIN for m in SOEID_RE.findall(str(cell))]


# --------------------------------------------------------------------------- #
# Loading the lookup tables
# --------------------------------------------------------------------------- #
def _xlsx_rows(path: Path) -> Iterator[Sequence]:
    """Rows of the first worksheet of an .xlsx/.xlsm file (cell values)."""
    from openpyxl import load_workbook   # only needed for Excel input
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from wb.active.iter_rows(values_only=True)
    finally:
        wb.close()


def _csv_rows(path: Path, logger: logging.Logger) -> Iterator[Sequence]:
    """Rows of a .csv file (all values are text).

    Encoding: UTF-8 (a BOM from Excel's "CSV UTF-8" is removed), falling back
    to Windows-1252 for files saved by Excel as plain "CSV".
    Delimiter: detected from the first lines (comma, semicolon, tab or pipe),
    comma if detection is not conclusive.
    """
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, newline="", encoding=encoding) as fh:
                fh.read()                     # decode the whole file once
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"{path}: cannot decode as {' or '.join(CSV_ENCODINGS)}")
    if encoding != CSV_ENCODINGS[0]:
        logger.warning(f"{path.name}: not UTF-8 - read as {encoding}")

    with open(path, newline="", encoding=encoding) as fh:
        sample = fh.read(64 * 1024)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        if delimiter != ",":
            logger.info(f"{path.name}: using '{delimiter}' as the CSV delimiter")
        for row in csv.reader(fh, delimiter=delimiter):
            if row and any(c.strip() for c in row):   # skip blank lines
                yield row


def segment_rows(path: Path, logger: logging.Logger) -> Iterator[Sequence]:
    """Rows of Managed_Segments as .xlsx/.xlsm or .csv, chosen by extension."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _xlsx_rows(path)
    if suffix in (".csv", ".txt"):
        return _csv_rows(path, logger)
    raise ValueError(f"{path}: unsupported segments file type '{path.suffix}' "
                     "(use .xlsx, .xlsm or .csv)")


def load_segments(path: Path, logger: logging.Logger
                  ) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
    """Read Managed_Segments (.xlsx or .csv) -> ({normalized_appid: {col: val}}, columns).

    `columns` is the file's column order with the key column removed.
    For Excel the first worksheet is used; in both formats the first row is
    the header.
    """
    rows = iter(segment_rows(path, logger))
    try:
        header = [squash_ws(h) for h in next(rows)]   # collapse header spacing
    except StopIteration:
        raise ValueError(f"{path} is empty")

    if SEGMENTS_KEY not in header:
        raise ValueError(
            f"{path}: key column '{SEGMENTS_KEY}' not found. Headers: {header}"
        )

    key_idx = header.index(SEGMENTS_KEY)
    data_cols = [c for c in header if c != SEGMENTS_KEY]

    lookup: Dict[str, Dict[str, str]] = {}
    dup = 0
    for raw in rows:
        if raw is None:
            continue
        cells = list(raw) + [None] * (len(header) - len(raw))  # pad short rows
        appid = normalize_appid(cells[key_idx])
        if not appid:
            continue
        record = {
            col: ("" if cells[i] is None else str(cells[i]).strip())
            for i, col in enumerate(header) if col != SEGMENTS_KEY
        }
        # Index under the normalized key AND the raw trimmed string, so an
        # EXACT Application Id matches whether the cell was stored as a number
        # or as text (preserving leading zeros). Matching is always full-value
        # equality - never a substring/prefix - so e.g. 123456 never matches
        # a 1234567 entry.
        raw_key = squash_ws(cells[key_idx])
        for key in {appid, raw_key} - {""}:
            if key in lookup:
                dup += 1
            lookup[key] = record  # last occurrence wins
    if hasattr(rows, "close"):
        rows.close()   # releases the workbook / file promptly

    if dup:
        logger.warning(f"Managed_Segments: {dup} duplicate Application Id(s); "
                       f"kept the last occurrence of each")
    logger.info(f"Managed_Segments: {len(lookup)} applications, "
                f"{len(data_cols)} detail columns")
    return lookup, data_cols


def load_inventory(path: Path, logger: logging.Logger
                   ) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
    """Read inventory.csv -> ({cluster: {col: val}}, columns).

    `columns` is the CSV column order with the key column removed.
    """
    with open(path, newline="", encoding="utf-8-sig") as fh:   # -sig: drop an Excel BOM
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            raise ValueError(f"{path} is empty")
        header = [squash_ws(h) for h in reader.fieldnames]  # collapse header spacing
        if INVENTORY_KEY not in header:
            raise ValueError(
                f"{path}: key column '{INVENTORY_KEY}' not found. "
                f"Headers: {header}"
            )
        data_cols = [c for c in header if c != INVENTORY_KEY]

        lookup: Dict[str, Dict[str, str]] = {}
        dup = 0
        for row in reader:
            # DictReader keys can carry irregular spacing if the header did;
            # rebuild against the space-collapsed header.
            clean = {squash_ws(h): (v.strip() if isinstance(v, str) else v)
                     for h, v in row.items() if h is not None}
            cluster = squash_ws(clean.get(INVENTORY_KEY))   # normalize the key
            if not cluster:
                continue
            if cluster in lookup:
                dup += 1
            lookup[cluster] = {c: clean.get(c, "") for c in data_cols}

    if dup:
        logger.warning(f"inventory: {dup} duplicate cluster(s); "
                       f"kept the last occurrence of each")
    logger.info(f"inventory: {len(lookup)} clusters, "
                f"{len(data_cols)} detail columns")
    return lookup, data_cols


# --------------------------------------------------------------------------- #
# Header assembly with collision handling
# --------------------------------------------------------------------------- #
def build_output_header(base_cols: List[str], seg_cols: List[str],
                        inv_cols: List[str]
                        ) -> Tuple[List[str], Dict[str, str], Dict[str, str]]:
    """Return (output_header, seg_rename, inv_rename).

    seg_rename / inv_rename map a source column name to the (possibly
    prefixed) name used in the output header, resolving collisions.
    """
    used = list(base_cols)
    used_set = set(base_cols)

    seg_rename: Dict[str, str] = {}
    for col in seg_cols:
        out = col if col not in used_set else f"{SEG_PREFIX}{col}"
        while out in used_set:                # extremely unlikely double clash
            out = f"{SEG_PREFIX}{out}"
        seg_rename[col] = out
        used.append(out)
        used_set.add(out)

    inv_rename: Dict[str, str] = {}
    for col in inv_cols:
        out = col if col not in used_set else f"{INV_PREFIX}{col}"
        while out in used_set:
            out = f"{INV_PREFIX}{out}"
        inv_rename[col] = out
        used.append(out)
        used_set.add(out)

    return used, seg_rename, inv_rename


# --------------------------------------------------------------------------- #
# Main build
# --------------------------------------------------------------------------- #
def build_notifybot(notification_path: Path, segments_path: Path,
                    inventory_path: Path, output_path: Path,
                    logger: logging.Logger) -> int:
    seg_lookup, seg_cols = load_segments(segments_path, logger)
    inv_lookup, inv_cols = load_inventory(inventory_path, logger)

    with open(notification_path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames is None:
            logger.error(f"{notification_path} is empty - nothing to build")
            return 1
        base_cols = [h.strip() for h in reader.fieldnames]

        if "appid" not in base_cols:
            logger.error(f"{notification_path}: expected an 'appid' column; "
                         f"headers are {base_cols}")
            return 1
        if "cluster" not in base_cols:
            logger.error(f"{notification_path}: expected a 'cluster' column; "
                         f"headers are {base_cols}")
            return 1

        # Output column order mirrors notification.csv, but the "ids" column
        # is renamed to "email".
        out_base_cols = [IDS_OUTPUT_COL if c == IDS_INPUT_COL else c
                         for c in base_cols]
        out_header, seg_rename, inv_rename = build_output_header(
            out_base_cols, seg_cols, inv_cols
        )

        # Blank templates so every row has every column.
        seg_blank = {seg_rename[c]: "" for c in seg_cols}
        inv_blank = {inv_rename[c]: "" for c in inv_cols}

        output_path.parent.mkdir(parents=True, exist_ok=True)

        n_rows = seg_hits = inv_hits = seg_miss = inv_miss = 0
        rows_enriched = emails_added = 0
        with open(output_path, "w", newline="") as out_fh:
            writer = csv.DictWriter(out_fh, fieldnames=out_header)
            writer.writeheader()

            for row in reader:
                record = {
                    (IDS_OUTPUT_COL if c == IDS_INPUT_COL else c): (row.get(c) or "")
                    for c in base_cols
                }

                appid_raw = squash_ws(row.get("appid"))
                appid = normalize_appid(row.get("appid"))
                seg_part = dict(seg_blank)
                seg = None
                if appid and appid.lower() not in EMPTY_APPIDS:
                    # Exact full-value match: try the raw appid, then its
                    # normalized form. No partial/substring matching.
                    seg = seg_lookup.get(appid_raw) or seg_lookup.get(appid)
                    if seg is not None:
                        for c in seg_cols:
                            seg_part[seg_rename[c]] = seg.get(c, "")
                        seg_hits += 1
                    else:
                        seg_miss += 1
                record.update(seg_part)

                cluster = squash_ws(row.get("cluster"))   # normalize the key
                inv_part = dict(inv_blank)
                inv = inv_lookup.get(cluster)
                if inv is not None:
                    for c in inv_cols:
                        inv_part[inv_rename[c]] = inv.get(c, "")
                    inv_hits += 1
                else:
                    inv_miss += 1
                record.update(inv_part)

                # Merge the owner / support-manager / app-manager SOEID emails
                # (from the matched segment) into the existing edit/view ids.
                soeid_emails: List[str] = []
                if seg is not None:
                    for col in SOEID_SOURCE_COLUMNS:
                        soeid_emails.extend(soeid_emails_from(seg.get(col, "")))
                if soeid_emails:
                    rows_enriched += 1
                    emails_added += len(soeid_emails)
                base_ids = [x.strip()
                            for x in (row.get(IDS_INPUT_COL) or "").split(",")
                            if x.strip()]
                seen, merged = set(), []
                for e in base_ids + soeid_emails:   # existing ids first
                    e = e.lower()                   # all ids in lower case
                    if e and e not in seen:
                        seen.add(e)
                        merged.append(e)
                joined = IDS_SEPARATOR.join(merged)
                if merged and IDS_TRAILING:
                    joined += IDS_SEPARATOR
                record[IDS_OUTPUT_COL] = joined

                writer.writerow(record)
                n_rows += 1

    logger.info("=" * 60)
    logger.info(f"notifybot rows written: {n_rows}")
    logger.info(f"  segment matches:  {seg_hits} hit / {seg_miss} miss")
    logger.info(f"  inventory matches:{inv_hits} hit / {inv_miss} miss")
    logger.info(f"  SOEID emails merged into ids: {emails_added} "
                f"across {rows_enriched} rows")
    logger.info(f"Output: {output_path}")
    logger.info("=" * 60)
    return 0 if n_rows > 0 else 1


def create_argument_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Join notification.csv with Managed_Segments.xlsx/.csv (by appid) "
                    "and inventory.csv (by cluster) into notifybot.csv",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("-n", "--notification", required=True, type=Path,
                   help="Path to notification.csv")
    p.add_argument("-s", "--segments", required=True, type=Path,
                   help="Path to Managed_Segments.xlsx or Managed_Segments.csv "
                        "(format chosen by extension)")
    p.add_argument("-i", "--inventory", required=True, type=Path,
                   help="Path to inventory.csv")
    p.add_argument("-o", "--output", type=Path,
                   help=f"Output file or folder (default: "
                        f"{DEFAULT_OUTPUT_DIR / OUTPUT_FILENAME})")
    p.add_argument("--log-level",
                   choices=["DEBUG", "INFO", "WARNING", "ERROR"], default="INFO")
    return p


def resolve_output(output_arg: Optional[Path]) -> Path:
    if output_arg is None:
        return DEFAULT_OUTPUT_DIR / OUTPUT_FILENAME
    if output_arg.is_dir() or output_arg.suffix == "":
        return output_arg / OUTPUT_FILENAME
    return output_arg


def main() -> int:
    wall_start = time.time()
    args = create_argument_parser().parse_args()
    logger = setup_logging(args.log_level)
    logger.info("NOTIFYBOT FLAT-FILE BUILD")

    for label, path in (("notification", args.notification),
                        ("segments", args.segments),
                        ("inventory", args.inventory)):
        if not path.exists():
            logger.error(f"{label} file not found: {path}")
            return 1

    output_path = resolve_output(args.output)

    try:
        rc = build_notifybot(args.notification, args.segments,
                             args.inventory, output_path, logger)
    except Exception as exc:
        import traceback
        logger.error(f"Build failed: {type(exc).__name__}: {exc}")
        logger.debug(traceback.format_exc())
        return 1

    logger.info(f"Elapsed: {time.time() - wall_start:.2f}s")
    if rc == 0:
        print(f"\nSUCCESS: notifybot.csv written -> {output_path}")
    else:
        print("\nNo rows written - check inputs")
    return rc


if __name__ == "__main__":
    sys.exit(main())
